# 심사위원 종이 채점표(.docx) 생성 스크립트.
#   python docs/make_judge_sheets.py  →  docs/심사위원_채점표.docx
#
# 심사위원은 웹이 아니라 이 종이로만 채점한다(0057·0059). 중간발표(9.28)·본선(10.9)
# 각각 [심사 기준 안내 1장 + 팀 목록이 미리 채워진 채점표]로 구성되며,
# 운영진은 회수한 채점표의 심사위원별 항목 합계(90점 만점)를 /admin/midterm ·
# /admin/scoring 표의 그 심사위원 칸에 옮겨 적는다 — 팀 평균은 사이트가 계산한다.
#
# 팀 목록은 .env.local 의 서비스 키로 DB(teams, 기권 제외, 조 번호순)에서 읽는다.
# 읽지 못하면(오프라인 등) 팀 이름 빈칸 40개로 만든다. 팀마다 세부 항목 표와
# 피드백 칸이 있는 블록 하나, 쪽당 2팀. 심사표 문구는 src/lib/types.ts 의
# FINAL_CRITERIA·MID_CRITERIA 와 같은 내용을 여기 직접 적는다 — 배점·항목이
# 바뀌면 여기와 DB(criteria) 둘 다 고칠 것.

import json
import sys
import urllib.request
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
NAVY = RGBColor(0x1F, 0x2A, 0x5C)
MUTED = RGBColor(0x66, 0x66, 0x66)
BLANK_ROWS = 40

# (항목, 배점, 설명, 심사위원 채점 여부, 세부 항목[(이름, 배점)])
# 세부 항목은 채점표에서 항목을 쪼개 적기 위한 것 — 배점 합이 항목 배점과 같아야 한다.
# 운영진은 항목 합계(90점)만 옮겨 적는다 — 세부 항목은 심사위원이 쓰는 칸이다.
MID = [
    ("논리의 연결성", 30, "문제 정의 → 해결 방안 → 기대 효과가 빈틈없이 이어지는지", True, [
        ("문제 정의가 명확하고 근거가 있는가", 10),
        ("해결 방안이 문제에서 자연스럽게 도출되는가", 10),
        ("기대 효과가 해결 방안과 논리적으로 이어지는가", 10),
    ]),
    ("실현 & 상용화 가능성", 20, "실제 환경에서 구현·활용될 수 있는지, 서비스·제품으로 발전할 가능성", True, [
        ("주어진 기간·기술로 구현할 수 있는가", 10),
        ("서비스·제품으로 발전·확장할 여지가 있는가", 10),
    ]),
    ("기획 문서의 재현 가능성", 20, "기획 문서만 보고도 같은 결과물을 만들 수 있을 만큼 구체적인지", True, [
        ("핵심 기능·요구사항이 구체적으로 정의되었는가", 10),
        ("화면·흐름·역할 분담·일정이 구체적인가", 10),
    ]),
    ("창의성 & 차별성", 20, "기존 서비스·해결방안 대비 독창성과 차별화된 특징", True, [
        ("아이디어·접근 방식이 독창적인가", 10),
        ("기존 서비스·해결방안과 뚜렷한 차별점이 있는가", 10),
    ]),
    ("발표", 10, "운영진이 반영 — 채점표에 없습니다", False, []),
]
FINAL = [
    ("실현 & 상용화 가능성", 30, "실제 환경에서 구현·활용될 수 있는지, 서비스·제품으로 발전할 가능성", True, [
        ("실제 지역 환경에서 바로 쓸 수 있는가", 10),
        ("운영·유지·확장이 현실적인가", 10),
        ("서비스·제품으로 발전할 가능성이 있는가", 10),
    ]),
    ("구현 완성도 & 기술력", 20, "핵심 기능의 실제 구현 여부와 활용 기술의 적절성·완성도", True, [
        ("핵심 기능이 실제로 동작하는가", 10),
        ("사용 기술이 적절하고 완성도가 높은가", 10),
    ]),
    ("지역 문제 적합성", 20, "지역사회 문제를 정확히 파악하고 그에 맞는 해결 방안을 제시했는지", True, [
        ("월계1동 지역 문제를 정확히 파악했는가", 10),
        ("해결 방안이 그 문제에 실제로 맞는가", 10),
    ]),
    ("창의성 & 차별성", 20, "기존 서비스·해결방안 대비 독창성과 차별화된 특징", True, [
        ("아이디어·접근 방식이 독창적인가", 10),
        ("기존 서비스·해결방안과 뚜렷한 차별점이 있는가", 10),
    ]),
    ("발표", 5, "운영진이 반영 — 채점표에 없습니다", False, []),
    ("참여도", 5, "운영진이 반영 — 개회식·최종발표 불참 시 인당 1점 감점", False, []),
]
for _rows in (MID, FINAL):
    for _name, _pts, _, _by_judge, _subs in _rows:
        assert not _by_judge or sum(p for _, p in _subs) == _pts, f"{_name}: 세부 배점 합 ≠ {_pts}"
TEAMS_PER_PAGE = 2

# 발표(채점) 순서 — 운영진이 정한 순서대로 채점표를 찍는다. 조 번호(team_no)순이 아니다.
# 공백을 뺀 이름으로 DB 팀과 맞춘다. 여기 없는 팀은 목록 맨 뒤에 조 번호순으로 붙는다.
PRESENT_ORDER = [
    "허강정", "바오밥나무", "NowonLikeUs", "지단", "월계디버깅",
    "월월계계", "MassCOM", "4 guys", "COMs", "입대 전 발악",
    "월계원정대", "월계방범대", "우럭아왜우럭", "얼마Geo", "솜사탕과 너구리",
    "이오", "월계상단", "월계 계섯거라", "떡잎마을방범대", "컴미컴",
    "CALAR", "라스트팡", "방과 후 지도타임", "로컬호스트", "노놀",
    "삼삼오오", "강컴퍼니", "월계동행", "월계획", "어쩌다 개발자",
    "이음(IEUM)", "배고프당", "유구무언", "탄탄대로", "일동차렷",
    "ESGenius", "복지나침반", "경영과컴퓨터", "ABC", "토큰좀주세요",
]
_ORDER_KEY = {n.replace(" ", ""): i for i, n in enumerate(PRESENT_ORDER)}


# ---------- 팀 목록 ----------
def load_env():
    env = {}
    p = ROOT / ".env.local"
    if not p.exists():
        return env
    for line in p.read_text(encoding="utf-8").splitlines():
        if "=" in line and not line.startswith("#"):
            k, v = line.split("=", 1)
            env[k.strip()] = v.strip().strip('"')
    return env


def fetch_teams():
    env = load_env()
    url, key = env.get("NEXT_PUBLIC_SUPABASE_URL"), env.get("SUPABASE_SERVICE_ROLE_KEY")
    if not url or not key:
        return None
    req = urllib.request.Request(
        f"{url}/rest/v1/teams?select=name,team_no&status=neq.withdrawn"
        "&order=team_no.asc.nullslast,name.asc",
        headers={"apikey": key, "Authorization": f"Bearer {key}"},
    )
    try:
        with urllib.request.urlopen(req, timeout=15) as res:
            # (조 번호, 이름). 번호가 없는 팀은 목록 뒤에 번호 없이 나온다(0058).
            teams = [(t.get("team_no"), t["name"]) for t in json.load(res)]
    except Exception as e:  # noqa: BLE001 — 실패해도 빈 표로 만든다
        print("팀 목록을 읽지 못해 빈 줄로 만듭니다:", e, file=sys.stderr)
        return None

    # 발표 순서로 정렬 — 목록에 없는 팀은 뒤에 (조 번호순으로) 붙이고 알려준다.
    missing = [n for _, n in teams if n.replace(" ", "") not in _ORDER_KEY]
    if missing:
        print("발표 순서 목록에 없어 맨 뒤로 보냅니다:", ", ".join(missing), file=sys.stderr)
    teams.sort(key=lambda t: _ORDER_KEY.get(t[1].replace(" ", ""), len(_ORDER_KEY)))
    return teams


# ---------- docx 도우미 ----------
def set_cell_bg(cell, hex_color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def no_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def row_height(row, cm):
    tr_pr = row._tr.get_or_add_trPr()
    h = OxmlElement("w:trHeight")
    h.set(qn("w:val"), str(int(cm * 567)))
    h.set(qn("w:hRule"), "atLeast")
    tr_pr.append(h)


def base_font(doc):
    style = doc.styles["Normal"]
    style.font.name = "맑은 고딕"
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    style.font.size = Pt(10.5)
    style.paragraph_format.space_after = Pt(4)
    for name in ("Heading 1", "Heading 2", "Title"):
        st = doc.styles[name]
        st.font.name = "맑은 고딕"
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
        st.font.color.rgb = NAVY


def portrait(section):
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    section.left_margin = section.right_margin = Cm(1.5)
    section.top_margin = section.bottom_margin = Cm(1.3)


def cell_text(cell, text, bold=False, size=None, color=None, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    if align:
        p.alignment = align
    r = p.add_run(text)
    r.bold = bold
    if size:
        r.font.size = Pt(size)
    if color:
        r.font.color.rgb = color
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    return r


def bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    if bold_prefix:
        r = p.add_run(bold_prefix)
        r.bold = True
    p.add_run(text)
    return p


def signature_line(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("심사위원 성명: ______________________    서명: ______________")
    r.font.size = Pt(11)


# ---------- 1장: 심사 기준 안내 ----------
def criteria_table(doc, rows):
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, text in zip(table.rows[0].cells, ("항목", "배점", "세부 항목 (배점)")):
        cell_text(cell, text, bold=True, size=10)
        set_cell_bg(cell, "E8ECF5")
    for name, pts, desc, by_judge, subs in rows:
        c = table.add_row().cells
        cell_text(c[0], name, bold=by_judge, size=10)
        cell_text(c[1], str(pts), size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
        if by_judge:
            cell_text(c[2], f"· {subs[0][0]} ({subs[0][1]})", size=9.5)
            for sub, sp in subs[1:]:
                para = c[2].add_paragraph()
                para.paragraph_format.space_after = Pt(0)
                r = para.add_run(f"· {sub} ({sp})")
                r.font.size = Pt(9.5)
        else:
            cell_text(c[2], desc, size=9.5)
            for cell in c:
                set_cell_bg(cell, "F5F5F5")
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.font.color.rgb = MUTED
    widths = (Cm(4.4), Cm(1.4), Cm(12.2))
    for row in table.rows:
        for cell, w in zip(row.cells, widths):
            cell.width = w


def guide_page(doc, title, when, intro, rows, notes, new_page=False):
    h1 = doc.add_heading("월계1동 지역문제 해결 해커톤", level=1)
    h1.paragraph_format.page_break_before = new_page
    h = doc.add_heading(f"{title} 심사위원 안내  ·  {when}", level=2)
    h.paragraph_format.space_before = Pt(0)
    doc.add_paragraph(intro)
    criteria_table(doc, rows)
    judge_items = [r for r in rows if r[3]]
    judge_total = sum(r[1] for r in judge_items)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    r = p.add_run(
        f"합계 100점. 심사위원께서는 회색 칸을 뺀 {len(judge_items)}개 항목, "
        f"{judge_total}점 만점으로 채점해 주세요. 발표·참여도는 운영진이 별도로 반영합니다."
    )
    r.bold = True
    doc.add_heading("채점 요령", level=2)
    for text, prefix in notes:
        bullet(doc, text, prefix)
    signature_line(doc)


# ---------- 채점표 ----------
def team_block(doc, seq, no, name, crits, judge_total):
    """팀 하나의 채점 블록 — 제목줄 · 세부 항목별 점수 · 합계 · 피드백 칸.

    제목줄은 발표 순번(seq)을 크게, 조 번호(no)는 작게 — 채점표는 발표 순서로
    찍지만 운영진은 조 번호로 점수를 옮겨 적기 때문에 둘 다 필요하다.
    """
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = (Cm(4.2), Cm(9.6), Cm(1.6), Cm(2.6))

    # 제목줄: No · 팀 이름 · 합계 적는 칸
    title = table.rows[0]
    row_height(title, 0.9)
    tc = title.cells
    left = tc[0].merge(tc[1])
    cell_text(left, f"{seq}.  ", bold=True, size=12)
    r = left.paragraphs[0].add_run(name if name else "팀: ______________________")
    r.bold = True
    r.font.size = Pt(12)
    if no:
        r2 = left.paragraphs[0].add_run(f"   ({no}조)")
        r2.font.size = Pt(9)
        r2.font.color.rgb = MUTED
    right = tc[2].merge(tc[3])
    cell_text(right, f"합계        / {judge_total}", bold=True, size=11, align=WD_ALIGN_PARAGRAPH.RIGHT)
    for cell in (left, right):
        set_cell_bg(cell, "E8ECF5")

    # 머리행
    hdr = table.add_row()
    for cell, text in zip(hdr.cells, ("항목", "세부 항목", "배점", "점수")):
        cell_text(cell, text, bold=True, size=9, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_bg(cell, "F5F5F5")

    # 세부 항목 — 같은 항목은 첫 칸을 세로로 합친다
    for cname, cpts, subs in crits:
        first = None
        for sub, sp in subs:
            row = table.add_row()
            row_height(row, 0.58)
            c = row.cells
            cell_text(c[1], sub, size=9.5)
            cell_text(c[2], str(sp), size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER)
            cell_text(c[3], "", size=10)
            if first is None:
                first = c[0]
            else:
                first = first.merge(c[0])
        cell_text(first, cname, bold=True, size=9.5)
        sub_r = first.paragraphs[0].add_run(f"  ({cpts})")
        sub_r.font.size = Pt(9)
        sub_r.font.color.rgb = MUTED

    # 피드백 칸 — 팀에 전달할 말·특이사항을 넉넉히 적을 수 있게 한 칸으로 크게
    fb = table.add_row()
    row_height(fb, 4.8)
    cell = fb.cells[0].merge(fb.cells[3])
    cell_text(cell, "피드백 · 코멘트", bold=True, size=9, color=MUTED)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP

    for row in table.rows:
        no_split(row)
        for cell, w in zip(row.cells, widths):
            cell.width = w


def score_sheet(doc, title, when, rows, teams):
    crits = [(n, p, subs) for n, p, _, j, subs in rows if j]
    judge_total = sum(p for _, p, _ in crits)
    names = teams if teams else [(None, "") for _ in range(BLANK_ROWS)]

    for i in range(0, len(names), TEAMS_PER_PAGE):
        # 쪽 나눔은 빈 단락이 아니라 제목의 page_break_before 로 — 표 뒤에 Word 가
        # 붙이는 빈 단락 + 나눔 단락이 겹치면 빈 쪽이 생긴다.
        h = doc.add_heading(f"{title} 채점표  ·  {when}", level=2)
        h.paragraph_format.page_break_before = True
        h.paragraph_format.space_before = Pt(0)
        h.paragraph_format.space_after = Pt(2)
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(
            f"세부 항목마다 0~배점 사이 정수로 적고 합계({judge_total}점 만점)를 계산해 주세요.    "
            "심사위원 성명: ______________"
        )
        r.font.size = Pt(9)
        r.font.color.rgb = MUTED
        for j, (no, name) in enumerate(names[i : i + TEAMS_PER_PAGE]):
            if j > 0:
                gap = doc.add_paragraph()
                gap.paragraph_format.space_after = Pt(2)
            team_block(doc, i + j + 1, no, name, crits, judge_total)


def main():
    teams = fetch_teams()
    doc = Document()
    base_font(doc)
    portrait(doc.sections[0])

    common_notes = [
        ("같은 항목이라도 팀 간 상대 비교를 염두에 두고 일관되게 매겨 주시면 집계가 공정해집니다.", "일관성 — "),
        ("발표 시간 준수·출석 같은 운영 요소는 운영진이 반영하므로 내용만 봐 주시면 됩니다.", "내용만 — "),
        ("점수와 순위는 대회가 끝날 때까지 비공개입니다. 채점표는 발표가 끝난 뒤 운영진이 회수합니다.", "비공개 — "),
        ("여러 심사위원의 점수는 팀별로 평균해 반영됩니다. 채점표에 성명·서명을 꼭 적어 주세요.", "합산 — "),
    ]

    # ---- 중간발표 ----
    guide_page(
        doc,
        "중간발표",
        "9.28(월) 18:30 · 80주년기념관 310호",
        "기획·아이디어 단계 발표입니다. 아래 4개 항목(90점)을 심사위원이 채점하고, "
        "발표 10점은 운영진이 반영해 합계 100점이 됩니다. 합산 1위 팀이 매니패스트상을 받습니다.",
        MID,
        [("중간발표 점수는 본선 점수에 합산되지 않습니다. 부담 없이 기획의 완성도로 평가해 주세요.", "본선과 별개 — ")]
        + common_notes,
    )
    score_sheet(doc, "중간발표", "9.28", MID, teams)

    # ---- 본선 ----
    guide_page(
        doc,
        "본선 최종발표",
        "10.9(금) 09:00 · 80주년기념관",
        "무박 2일 개발을 마친 결과물 발표입니다(팀당 5분 + 질의응답). 아래 4개 항목(90점)을 심사위원이 채점하고, "
        "발표 5점·참여도 5점은 운영진이 반영해 합계 100점이 됩니다.",
        FINAL,
        [
            ("심사위원 점수 : 참가 팀 상호평가 = 2 : 1 로 합산해 상위 15팀을 전시 진출팀으로 뽑습니다.", "1차 선정 — "),
            ("진출팀은 전시 기간(10.11–10.13) 주민투표 점수를 더합니다. 심사 50% · 상호평가 25% · 주민투표 25%. "
             "합산 1위 노원구청장상, 2~4위 광운대학교 총장상.", "최종 순위 — "),
        ]
        + common_notes,
        new_page=True,
    )
    score_sheet(doc, "본선 최종발표", "10.9", FINAL, teams)

    out = ROOT / "docs" / "심사위원_채점표.docx"
    doc.save(out)
    print(out, f"(팀 {len(teams) if teams else BLANK_ROWS}개, 쪽당 {TEAMS_PER_PAGE}팀)")


if __name__ == "__main__":
    main()
