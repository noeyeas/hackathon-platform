# 심사위원 안내서(.docx) 생성 스크립트.
#   python docs/make_judge_guide.py  →  docs/심사위원_안내.docx
# 심사표 문구는 src/lib/types.ts 의 FINAL_CRITERIA·MID_CRITERIA 와 같은 내용을
# 여기 직접 적는다(문서는 사이트와 별개로 배포되므로 DB 를 읽지 않는다).
# 배점·항목이 바뀌면 여기와 DB(criteria) 둘 다 고칠 것.

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from pathlib import Path

SITE = "https://www.kw-hackathon.co.kr"
NAVY = RGBColor(0x1F, 0x2A, 0x5C)
MUTED = RGBColor(0x66, 0x66, 0x66)

MID = [
    ("논리의 연결성", 30, "문제 정의 → 해결 방안 → 기대 효과가 빈틈없이 이어지는지"),
    ("실현 & 상용화 가능성", 20, "실제 환경에서 구현·활용될 수 있는지, 서비스·제품으로 발전할 가능성"),
    ("기획 문서의 재현 가능성", 20, "기획 문서만 보고도 같은 결과물을 만들 수 있을 만큼 구체적인지"),
    ("창의성 & 차별성", 20, "기존 서비스·해결방안 대비 독창성과 차별화된 특징"),
    ("발표", 10, "운영진이 반영 — 심사위원 채점표에는 없습니다"),
]
FINAL = [
    ("실현 & 상용화 가능성", 30, "실제 환경에서 구현·활용될 수 있는지, 서비스·제품으로 발전할 가능성"),
    ("구현 완성도 & 기술력", 20, "핵심 기능의 실제 구현 여부와 활용 기술의 적절성·완성도"),
    ("지역 문제 적합성", 20, "지역사회 문제를 정확히 파악하고 그에 맞는 해결 방안을 제시했는지"),
    ("창의성 & 차별성", 20, "기존 서비스·해결방안 대비 독창성과 차별화된 특징"),
    ("발표", 5, "운영진이 반영 — 심사위원 채점표에는 없습니다"),
    ("참여도", 5, "운영진이 반영 — 개회식·최종발표 불참 시 인당 1점 감점"),
]


def set_cell_bg(cell, hex_color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def base_font(doc):
    style = doc.styles["Normal"]
    style.font.name = "맑은 고딕"
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    style.font.size = Pt(10.5)
    for name in ("Heading 1", "Heading 2", "Title"):
        st = doc.styles[name]
        st.font.name = "맑은 고딕"
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
        st.font.color.rgb = NAVY


def h1(doc, text):
    p = doc.add_heading(text, level=1)
    p.paragraph_format.space_before = Pt(14)
    return p


def para(doc, text="", bold_prefix=None, muted=False):
    p = doc.add_paragraph()
    if bold_prefix:
        r = p.add_run(bold_prefix)
        r.bold = True
    r = p.add_run(text)
    if muted:
        r.font.color.rgb = MUTED
        r.font.size = Pt(9.5)
    return p


def bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    if bold_prefix:
        r = p.add_run(bold_prefix)
        r.bold = True
    p.add_run(text)
    return p


def numbered(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style="List Number")
    if bold_prefix:
        r = p.add_run(bold_prefix)
        r.bold = True
    p.add_run(text)
    return p


def criteria_table(doc, rows, judge_note):
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for cell, text in zip(hdr, ("항목", "배점", "설명")):
        cell.text = ""
        run = cell.paragraphs[0].add_run(text)
        run.bold = True
        set_cell_bg(cell, "E8ECF5")
    for name, pts, desc in rows:
        c = table.add_row().cells
        c[0].text = name
        c[1].text = str(pts)
        c[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        c[2].text = desc
        if "운영진이 반영" in desc:
            for cell in c:
                set_cell_bg(cell, "F5F5F5")
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.font.color.rgb = MUTED
    widths = (Cm(4.2), Cm(1.4), Cm(10.4))
    for row in table.rows:
        for cell, w in zip(row.cells, widths):
            cell.width = w
    total = sum(p for _, p, _ in rows)
    para(doc, f"합계 {total}점. {judge_note}", muted=True)


def build():
    doc = Document()
    base_font(doc)
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Cm(2)
        s.left_margin = s.right_margin = Cm(2.2)

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("월계1동 지역문제 해결 해커톤")
    r.font.size = Pt(11)
    r.font.color.rgb = MUTED
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("심사위원 안내서")
    r.bold = True
    r.font.size = Pt(22)
    r.font.color.rgb = NAVY
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run(f"채점은 {SITE} 에서 휴대폰·노트북으로 진행합니다. 5분이면 다 읽을 수 있습니다.")
    r.font.color.rgb = MUTED

    # ── 한눈에
    h1(doc, "한눈에 보기")
    numbered(doc, "로 로그인합니다 (Google 로그인 권장).", "등록된 이메일")
    numbered(doc, " 메뉴 → 날짜에 맞는 탭을 고릅니다: 중간발표 (9.28) / 본선 최종발표 (10.9).", "상단 ‘심사’")
    numbered(doc, "를 눌러 펼치고, 항목마다 슬라이더로 점수를 정한 뒤 ‘점수 저장’을 누릅니다.", "팀 카드")
    numbered(doc, "이 ‘모두 완료’가 될 때까지 반복합니다. 저장한 점수는 언제든 다시 열어 고칠 수 있습니다.", "상단 진행률")

    # ── 1. 로그인
    h1(doc, "1. 로그인")
    para(doc, "운영진에 알려 주신 이메일이 심사위원으로 미리 등록되어 있습니다. 그 이메일로 로그인하면 자동으로 심사위원 권한이 부여됩니다.")
    bullet(doc, "을 권장합니다 — 즉시 됩니다.", "Google 로그인")
    bullet(doc, "이메일 링크 로그인도 되지만 메일 도착까지 잠시 걸릴 수 있습니다.")
    bullet(doc, "(예: 학교 계정 ↔ Gmail)으로 로그인하면 일반 참가자로 인식돼 채점 화면이 열리지 않습니다. 등록된 주소인지 확인해 주세요.", "다른 계정")
    para(doc, f"채점 화면 주소: {SITE}/judge  (중간발표: {SITE}/judge/mid)", muted=True)

    # ── 2. 화면
    h1(doc, "2. 채점 화면 사용법")
    bullet(doc, " — 중간발표와 본선은 심사표가 다르므로 화면 위에서 날짜에 맞는 탭을 고릅니다. 잘못된 탭에서는 저장이 되지 않으니 걱정하지 않으셔도 됩니다.", "탭")
    bullet(doc, " — 화면 위 막대가 ‘채점 완료 n / 전체 팀’을 보여 줍니다. 남은 팀 수가 0이 되면 끝입니다.", "진행률")
    bullet(doc, " — 팀 이름을 누르면 채점표가 펼쳐집니다. 오른쪽 배지가 ‘미채점’이면 저장 전, ‘✓ 채점됨’이면 저장된 상태입니다.", "팀 카드")
    bullet(doc, " — 항목마다 0점부터 배점까지 1점 단위로 움직입니다. 오른쪽 숫자가 ‘—’이면 아직 손대지 않은 것입니다. 0점을 주려면 슬라이더를 한 번 눌러 주세요.", "슬라이더")
    bullet(doc, " — 카드 아래에 항목 합계가 자동으로 표시됩니다.", "합계")
    bullet(doc, " (선택) — 팀에 남기고 싶은 말이 있으면 적어 주세요. 운영진만 볼 수 있습니다.", "코멘트")
    bullet(doc, " — 모든 항목을 채워야 저장됩니다. 저장 후에도 카드를 다시 열어 수정할 수 있습니다.", "점수 저장")
    bullet(doc, "저장이 안 되고 ‘평가가 닫혀 있습니다’가 보이면 운영진이 아직 채점을 열지 않은 것입니다. 행사 진행에 맞춰 열립니다.")

    # ── 3. 중간
    h1(doc, "3. 중간발표 (9.28 18:30, 80주년기념관 310호)")
    para(doc, "기획·아이디어 단계 발표입니다. 아래 4개 항목(90점)을 심사위원이 채점하고, 발표 10점은 운영진이 별도로 반영해 합계 100점이 됩니다. 합산 1위 팀이 매니패스트상을 받습니다.")
    criteria_table(doc, MID, "심사위원 채점표는 회색 칸을 뺀 4항목 90점입니다.")
    para(doc, "중간발표 점수는 본선 점수에 합산되지 않습니다. 부담 없이 기획의 완성도로 평가해 주세요.", muted=True)

    # ── 4. 본선
    h1(doc, "4. 본선 최종발표 (10.9 09:00, 80주년기념관)")
    para(doc, "무박 2일 개발을 마친 결과물 발표입니다(팀당 5분 + 질의응답). 아래 4개 항목(90점)을 심사위원이 채점하고, 발표 5점·참여도 5점은 운영진이 반영해 합계 100점이 됩니다.")
    criteria_table(doc, FINAL, "심사위원 채점표는 회색 칸을 뺀 4항목 90점입니다.")
    para(doc, "", bold_prefix="본선 점수는 이렇게 합쳐집니다")
    bullet(doc, " — 심사위원 점수 : 참가 팀 상호평가 = 2 : 1 로 합산해 상위 15팀을 전시 진출팀으로 뽑습니다.", "1차 선정")
    bullet(doc, " — 진출팀은 여기에 전시 기간(10.11–10.13) 주민투표 점수를 더합니다. 심사 50% · 상호평가 25% · 주민투표 25%.", "최종 순위")
    bullet(doc, "합산 1위 노원구청장상, 2~4위 광운대학교 총장상.")
    para(doc, "심사위원 점수가 가장 큰 비중(50%)입니다. 여러 심사위원의 점수는 100점 만점으로 환산한 뒤 평균됩니다.", muted=True)

    # ── 5. 유의
    h1(doc, "5. 유의 사항")
    bullet(doc, "점수와 순위는 대회가 끝날 때까지 비공개입니다. 참가자는 심사위원 개개인의 점수를 볼 수 없습니다.")
    bullet(doc, "같은 항목이라도 팀 간 상대 비교를 염두에 두고 일관되게 매겨 주시면 집계가 공정해집니다. 필요하면 나중에 되돌아가 수정하셔도 됩니다.")
    bullet(doc, "발표 시간 준수·출석 같은 운영 요소는 운영진이 반영하므로 심사위원께서는 내용만 봐 주시면 됩니다.")
    bullet(doc, "기권한 팀은 목록에 나오지 않습니다.")
    bullet(doc, "화면이 이상하거나 저장이 안 되면 새로고침 후 다시 시도하고, 그래도 안 되면 현장 운영진에게 말씀해 주세요.")

    out = Path(__file__).with_name("심사위원_안내.docx")
    doc.save(out)
    print(out)


if __name__ == "__main__":
    build()
